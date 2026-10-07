#!/usr/bin/env python3
"""专业版微信头像评估报告生成模板（python-docx）

用法：
  1. 复制本脚本为 generate_{姓名拼音}_doc.py
  2. 修改下方 CONFIG 区的变量
  3. 按各章节 TODO 填入命主具体分析内容
  4. 运行：python3 generate_xxx_doc.py

依赖：python-docx（已装于 envs/default venv）
"""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import os

# ============================================================
# CONFIG — 生成前修改此处
# ============================================================
NAME = "案例姓名"                      # 命主姓名
IMG_PATH = "/path/to/avatar.jpg"       # 头像图片路径（可为 None）
OUT_PATH = os.path.expanduser(f"~/Desktop/微信头像调整/{NAME}-微信头像评估报告.docx")

# 配色
C_PRIMARY = RGBColor(0x8B, 0x45, 0x13)
C_ACCENT  = RGBColor(0xC0, 0x39, 0x2B)
C_SOFT    = RGBColor(0xD4, 0xA5, 0x76)
C_TEXT    = RGBColor(0x3E, 0x3E, 0x3E)
C_LIGHT   = RGBColor(0x7A, 0x6A, 0x5A)
C_GREEN   = RGBColor(0x4A, 0x7C, 0x59)
C_BLUE    = RGBColor(0x33, 0x66, 0x99)

BG_SOFT, BG_WARM, BG_GREEN = "FAF3E8", "FFF5EE", "F0F7F4"
BG_GOLD, BG_BLUE = "FDF6E3", "E8F0F8"

doc = Document()

style = doc.styles['Normal']
style.font.name = '微软雅黑'
style.font.size = Pt(11)
style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
style.paragraph_format.line_spacing = 1.5

for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

# ============================================================
# 工具函数
# ============================================================

def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}" w:val="clear"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def heading(text, level=1, color=C_PRIMARY):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = color
        r.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
    return h


def para(text, bold=False, color=None, align=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.font.name = '微软雅黑'
    run.font.size = Pt(11)
    run.font.bold = bold
    run.font.color.rgb = color or C_TEXT
    run.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
    p.paragraph_format.line_spacing = 1.5
    return p


def add_bullet(text, doc_obj=None):
    d = doc_obj or doc
    p = d.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.name = '微软雅黑'
    run.font.size = Pt(10.8)
    run.font.color.rgb = C_TEXT
    run.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
    return p


def add_picture(path, width=Inches(2.6)):
    if not path:
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    try:
        p.add_run().add_picture(path, width=width)
    except Exception as e:
        print(f"Warning: could not insert image: {e}")
    p.paragraph_format.space_after = Pt(6)


def card(title, items, bg_color=BG_SOFT, title_color=C_PRIMARY):
    """卡片式文本块（单元格表格模拟）"""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, bg_color)
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:color="D4A576"/>'
        f'<w:bottom w:val="single" w:sz="4" w:color="D4A576"/>'
        f'<w:left w:val="single" w:sz="4" w:color="D4A576"/>'
        f'<w:right w:val="single" w:sz="4" w:color="D4A576"/>'
        f'</w:tcBorders>'
    ))
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(8)
    run = p.add_run(title)
    run.font.name = '微软雅黑'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = title_color
    run.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
    for item in items:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(f'  {item}')
        run.font.name = '微软雅黑'
        run.font.size = Pt(10.5)
        run.font.color.rgb = C_TEXT
        run.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
    doc.add_paragraph()


def data_table(header, rows, col_widths=None):
    """标准数据表：蓝底白字表头 + 交替行底色"""
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    # 表头
    for j, h in enumerate(header):
        cell = tbl.cell(0, j)
        cell.text = ''
        run = cell.paragraphs[0].add_run(h)
        run.font.name = '微软雅黑'
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, "2F5496")
    # 数据行（交替底色）
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i + 1, j)
            cell.text = ''
            run = cell.paragraphs[0].add_run(str(val))
            run.font.name = '微软雅黑'
            run.font.size = Pt(10.5)
            run.font.color.rgb = C_TEXT
            run.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
            if i % 2 == 1:
                set_cell_shading(cell, "EAF1F8")
    if col_widths:
        for j, w in enumerate(col_widths):
            for row in tbl.rows:
                row.cells[j].width = Cm(w)
    doc.add_paragraph()
    return tbl


# ============================================================
# 封面
# ============================================================
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(20)
run = p.add_run('微信头像评估报告')
run.font.name = '微软雅黑'
run.font.size = Pt(22)
run.font.bold = True
run.font.color.rgb = C_PRIMARY
run.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(16)
run = p.add_run(f'——  {NAME}  ——')
run.font.name = '微软雅黑'
run.font.size = Pt(12)
run.font.color.rgb = C_LIGHT
run.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run(f'报告日期：__年__月__日')
run.font.size = Pt(10)
run.font.color.rgb = C_LIGHT

# ============================================================
# 一、八字排盘与命格分析
# ============================================================
heading('一、八字排盘与命格分析')
para('TODO：列出四柱八字（年柱 月柱 日柱 时柱）、日主、纳音五行。')
# 示例：
# data_table(['柱', '年柱', '月柱', '日柱', '时柱'],
#            [['干支', '甲子', '乙亥', '乙卯', '丙戌']])
para('TODO：五行旺衰简析（得令/得地/得势、身旺身弱判断）。')
para('TODO：用神与忌神结论。')

# ============================================================
# 二、头像五行拆解
# ============================================================
heading('二、头像五行拆解')
add_picture(IMG_PATH)
para('TODO：逐一列出画面元素 → 五行归属 → 吉凶判断（用 data_table）。')
para('TODO：五行占比总结（如：木水偏旺、火土偏弱）。')

# ============================================================
# 三、四象限分析
# ============================================================
heading('三、四象限分析')
data_table(
    ['方位', '四象', '本位五行', '头像呈现', '吉凶简评'],
    [
        # TODO: 填入 上方/下方/左方/右方/中央 五行分析
        ['上方', '朱雀', '火', 'TODO', 'TODO'],
        ['下方', '玄武', '水', 'TODO', 'TODO'],
        ['左方', '青龙', '木', 'TODO', 'TODO'],
        ['右方', '白虎', '金', 'TODO', 'TODO'],
        ['中央', '太极', '土', 'TODO', 'TODO'],
    ],
    col_widths=[2.5, 2, 2.5, 4.5, 4.5],
)
para('TODO：天地阴阳判断（人物正向无颠倒？阳气是否充沛？）。')

# ============================================================
# 四、健康分析
# ============================================================
heading('四、健康分析')
para('TODO：按命局+头像双重叠加，引用克伐疾患速查（木旺克土→脾胃；火旺克金→肺肠；'
     '土旺克水→肾膀；金旺克木→肝胆；水旺克火→心小肠）。')
para('TODO：五脏开窍对应（肾-耳、肝-目、脾-口、肺-鼻、心-舌）。')

# ============================================================
# 五、财位分析
# ============================================================
heading('五、财位分析')
data_table(
    ['财位类型', '对应方位/五行', '头像现状', '评估'],
    [
        # TODO: 年命财位（纳音）/ 日命财位（日主）/ 时命财位 / 八宅财位
        ['年命财位', 'TODO', 'TODO', 'TODO'],
        ['日命财位', 'TODO', 'TODO', 'TODO'],
        ['时命财位', 'TODO', 'TODO', 'TODO'],
        ['八宅财位', 'TODO', 'TODO', 'TODO'],
    ],
    col_widths=[3.5, 4, 5, 2.5],
)
para('TODO：财位综合判断。')

# ============================================================
# 六、综合评分
# ============================================================
heading('六、综合评分')
data_table(
    ['维度', '得分', '说明'],
    [
        ['太极点清晰度', 'TODO/5', 'TODO'],
        ['四象限五行平衡', 'TODO/5', 'TODO'],
        ['用神补益度', 'TODO/5', 'TODO'],
        ['情感与人物特质', 'TODO/5', 'TODO'],
        ['健康暗示', 'TODO/5', 'TODO'],
        ['总分', 'TODO/25', '折算约 TODO/10'],
    ],
    col_widths=[4.5, 2.5, 8],
)

# ============================================================
# 七、调整建议
# ============================================================
heading('七、调整建议')
para('TODO：至少 2–3 套方案，每套含思路 + 具体做法（服装颜色/背景/饰品/避免元素）。')

# ============================================================
# 八、养生建议
# ============================================================
heading('八、养生建议')
para('TODO：饮食调养 / 作息运动 / 情志调摄 三方面，结合命主过旺过弱五行给出。')

# ============================================================
# 九、结语
# ============================================================
heading('九、结语')
para('TODO：总结核心问题与调整方向，附祝福语（如：愿您五行调和，身心暖阳，财源稳固，喜乐安康）。')

os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
doc.save(OUT_PATH)
print(f"专业版报告已生成：{OUT_PATH}")
