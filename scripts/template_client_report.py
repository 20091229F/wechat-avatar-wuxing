#!/usr/bin/env python3
"""客户版微信头像报告生成模板（python-docx 温馨简明版）

用法：
  1. 复制本脚本为 generate_{姓名拼音}_client.py
  2. 修改 CONFIG 区变量
  3. 按各章节 TODO 填入客户化内容（生活化语言，少专业术语）
  4. 运行：python3 generate_xxx_client.py

依赖：python-docx
"""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import os

# ============================================================
# CONFIG — 生成前修改此处
# ============================================================
NAME = "客户姓名"
WUXING_METAPHOR = "乙木（花草之木）"   # 五行比喻，如：戊土（城墙之土）/ 壬水（江河之水）
IMG_PATH = "/path/to/avatar.jpg"
OUT_PATH = os.path.expanduser(f"~/Desktop/微信头像调整/{NAME}-客户版报告.docx")

C_PRIMARY = RGBColor(0x8B, 0x45, 0x13)
C_ACCENT  = RGBColor(0xC0, 0x39, 0x2B)
C_SOFT    = RGBColor(0xD4, 0xA5, 0x76)
C_TEXT    = RGBColor(0x3E, 0x3E, 0x3E)
C_LIGHT   = RGBColor(0x7A, 0x6A, 0x5A)
C_GREEN   = RGBColor(0x4A, 0x7C, 0x59)
C_BLUE    = RGBColor(0x33, 0x66, 0x99)

BG_SOFT, BG_WARM, BG_GREEN = "FAF3E8", "FFF5EE", "F0F7F4"
BG_GOLD, BG_BLUE = "FDF6E3", "E8F0F8"

doc = Document()

style = doc.styles['Normal']
style.font.name = '微软雅黑'
style.font.size = Pt(11)
style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
style.paragraph_format.line_spacing = 1.5

for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)


def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}" w:val="clear"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def heading(text, level=1, color=C_PRIMARY):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = color
        r.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
    return h


def para(text, bold=False, color=None, align=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.font.name = '微软雅黑'
    run.font.size = Pt(11)
    run.font.bold = bold
    run.font.color.rgb = color or C_TEXT
    run.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
    p.paragraph_format.line_spacing = 1.5
    return p


def add_picture(path, width=Inches(2.6)):
    if not path:
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    try:
        p.add_run().add_picture(path, width=width)
    except Exception as e:
        print(f"Warning: could not insert image: {e}")
    p.paragraph_format.space_after = Pt(6)


def card(title, items, bg_color=BG_SOFT, title_color=C_PRIMARY):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, bg_color)
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:color="D4A576"/>'
        f'<w:bottom w:val="single" w:sz="4" w:color="D4A576"/>'
        f'<w:left w:val="single" w:sz="4" w:color="D4A576"/>'
        f'<w:right w:val="single" w:sz="4" w:color="D4A576"/>'
        f'</w:tcBorders>'
    ))
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(8)
    run = p.add_run(title)
    run.font.name = '微软雅黑'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = title_color
    run.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
    for item in items:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(f'  {item}')
        run.font.name = '微软雅黑'
        run.font.size = Pt(10.5)
        run.font.color.rgb = C_TEXT
        run.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
    doc.add_paragraph()


# ============================================================
# 封面
# ============================================================
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(20)
run = p.add_run('微信头像专属分析报告')
run.font.name = '微软雅黑'
run.font.size = Pt(22)
run.font.bold = True
run.font.color.rgb = C_PRIMARY
run.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(16)
run = p.add_run(f'——  {NAME}  ——')
run.font.name = '微软雅黑'
run.font.size = Pt(12)
run.font.color.rgb = C_LIGHT
run.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
run = p.add_run('每一张头像，都在悄悄影响着您的人缘与运势。')
run.font.size = Pt(11)
run.font.italic = True
run.font.color.rgb = C_LIGHT

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(16)
run = p.add_run('愿这份报告，帮您找到最适合自己的那一张。')
run.font.size = Pt(11)
run.font.italic = True
run.font.color.rgb = C_LIGHT

# ============================================================
# 一、关于您
# ============================================================
heading('一、关于您')
# TODO: 用五行比喻介绍命主特质，生活化语言
para(f'您属于「{WUXING_METAPHOR}」日主——TODO：展开该五行的性格比喻。')
para('TODO：通俗解释命局旺衰（如"像一株被浇了太多水的花""像一片被大雪覆盖的土地"），'
     '以及最需要什么五行来平衡。')

card('您的五行特质', [
    '日主：TODO（比喻）',
    '特质：TODO',
    '需要：TODO（温暖/扎根/滋养等生活化表达）',
    '注意：TODO',
], bg_color=BG_GOLD, title_color=C_PRIMARY)

# ============================================================
# 二、当前头像解读
# ============================================================
heading('二、当前头像为什么需要调整')
para('TODO：先真诚夸赞画面气质（先扬后抑），再用生活化语言说明问题。')
add_picture(IMG_PATH)
# TODO: bullet 式问题列表，每条用比喻解释

# ============================================================
# 三、更适合您的头像方向
# ============================================================
heading('三、更适合您的头像方向')
# TODO: 2-3 套方案卡片，每套含服装颜色/背景/配饰/避免元素
card('推荐方案一：TODO（最推荐）', [
    'TODO',
], bg_color=BG_WARM, title_color=C_ACCENT)
card('推荐方案二：TODO', ['TODO'], bg_color=BG_BLUE, title_color=C_BLUE)
card('推荐方案三：TODO', ['TODO'], bg_color=BG_GREEN, title_color=C_GREEN)

# ============================================================
# 四、用色建议
# ============================================================
heading('四、用色建议')
card('宜用色', [
    'TODO：补X补Y，最利您',
], bg_color=BG_SOFT, title_color=C_PRIMARY)
card('忌用色', [
    'TODO：原因的生活化解释',
], bg_color=BG_WARM, title_color=C_ACCENT)

# ============================================================
# 五、健康小贴士
# ============================================================
heading('五、健康小贴士')
card('需重点关注的身体部位', [
    'TODO——TODO',
], bg_color=BG_GOLD, title_color=C_PRIMARY)
para('日常调养小建议：')
for tip in ['TODO：饮食', 'TODO：作息', 'TODO：运动', 'TODO：情志']:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    run = p.add_run(f'  {tip}')
    run.font.size = Pt(10.8)
    run.font.color.rgb = C_TEXT
    run.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')

# ============================================================
# 六、温馨祝福
# ============================================================
heading('温馨祝福', color=C_PRIMARY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(10)
run = p.add_run(f'亲爱的{NAME}，')
run.font.size = Pt(12)
run.font.color.rgb = C_PRIMARY
run.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')

# TODO: 3-4 段情感共鸣文字：感谢信任 → 换头像的意义 → 换后向世界传递的信号 → 五行祝福
para('TODO', align=WD_ALIGN_PARAGRAPH.CENTER)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(12)
run = p.add_run('愿您万事顺遂，福慧双修')
run.font.size = Pt(13)
run.font.bold = True
run.font.color.rgb = C_ACCENT
run.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')

# 报告信息
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('报告日期：__年__月__日')
run.font.size = Pt(9)
run.font.color.rgb = C_LIGHT

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('——  本报告为您专属定制，请妥善保管  ——')
run.font.size = Pt(9)
run.font.color.rgb = C_LIGHT

os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
doc.save(OUT_PATH)
print(f"客户版报告已生成：{OUT_PATH}")
